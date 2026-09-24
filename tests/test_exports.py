from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader


def test_dashboard_excel_contains_kpis(client):
    client.post("/analyze", json={"text": "Robot robot science python."})
    response = client.get("/api/dashboard/export/excel")
    assert response.status_code == 200
    assert response.data[:2] == b"PK"
    workbook = load_workbook(BytesIO(response.data))
    assert "KPIs" in workbook.sheetnames
    assert workbook["KPIs"]["A1"].value
    workbook.close()


def test_dashboard_pdf_is_readable(client):
    response = client.get("/api/dashboard/export/pdf")
    assert response.data.startswith(b"%PDF")
    reader = PdfReader(BytesIO(response.data))
    assert len(reader.pages) >= 1
    assert "YouTube Transcript Analyzer" in " ".join(reader.pages[0].extract_text().split())

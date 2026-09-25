"""
Report Export Service - Phase 6
Service pour générer des rapports PDF/Excel avec les métriques du dashboard
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .dashboard_service import dashboard_service

logger = logging.getLogger(__name__)


class ReportExportService:
    def __init__(self) -> None:
        self.styles = getSampleStyleSheet()
        self._create_custom_styles()

    def _create_custom_styles(self) -> None:
        """Créer des styles personnalisés pour les rapports"""
        self.styles.add(
            ParagraphStyle(
                name="CustomTitle",
                parent=self.styles["Heading1"],
                fontSize=24,
                spaceAfter=30,
                textColor=colors.HexColor("#667eea"),
                alignment=1,  # Center
            )
        )

        self.styles.add(
            ParagraphStyle(
                name="CustomHeading2",
                parent=self.styles["Heading2"],
                fontSize=16,
                spaceBefore=20,
                spaceAfter=10,
                textColor=colors.HexColor("#2c3e50"),
            )
        )

        self.styles.add(
            ParagraphStyle(
                name="KPIValue",
                parent=self.styles["Normal"],
                fontSize=18,
                textColor=colors.HexColor("#667eea"),
                alignment=1,  # Center
                spaceAfter=10,
            )
        )

    def generate_dashboard_report(self, days: int = 30, format_type: str = "pdf") -> io.BytesIO:
        """
        Génère un rapport complet du dashboard
        """
        try:
            # Récupérer les données du dashboard
            dashboard_data = dashboard_service.generate_comprehensive_dashboard(days=days)

            if not dashboard_data:
                raise ValueError("Impossible de récupérer les données du dashboard")

            if format_type.lower() == "pdf":
                return self._generate_pdf_report(dashboard_data, days)
            elif format_type.lower() == "excel":
                return self._generate_excel_report(dashboard_data, days)
            else:
                raise ValueError(f"Format non supporté: {format_type}")

        except Exception as e:
            logger.error(f"Erreur lors de la génération du rapport: {e}")
            raise

    def _generate_pdf_report(self, data: dict[str, Any], days: int) -> io.BytesIO:
        """Génère un rapport PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18
        )

        # Construire le contenu du rapport
        story = []

        # Titre principal
        title = Paragraph(
            "📊 Rapport Analytics YouTube Transcript Analyzer", self.styles["CustomTitle"]
        )
        story.append(title)
        story.append(Spacer(1, 12))

        # Informations du rapport
        period_info = f"Période analysée: {days} derniers jours<br/>"
        period_info += f"Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M')}<br/>"
        if data.get("kpis"):
            period_info += f"Du {data['kpis']['period']['start_date'][:10]} au {data['kpis']['period']['end_date'][:10]}"

        story.append(Paragraph(period_info, self.styles["Normal"]))
        story.append(Spacer(1, 20))

        # Section KPIs
        story.extend(self._create_kpis_section(data.get("kpis", {})))

        # Section Tendances
        story.extend(self._create_trends_section(data.get("trends", {})))

        # Section Insights
        story.extend(self._create_insights_section(data.get("insights", {})))

        # Construire le PDF
        doc.build(story)
        buffer.seek(0)
        return buffer

    def _create_kpis_section(self, kpis: dict[str, Any]) -> list:
        """Crée la section KPIs du rapport"""
        elements = []

        elements.append(
            Paragraph("📈 Indicateurs Clés de Performance (KPIs)", self.styles["CustomHeading2"])
        )
        elements.append(Spacer(1, 12))

        if not kpis or not kpis.get("totals"):
            elements.append(Paragraph("Aucune donnée KPI disponible", self.styles["Normal"]))
            return elements

        # Tableau KPIs
        kpi_data = [
            ["Métrique", "Valeur", "Description"],
            [
                "Analyses (période)",
                f"{kpis['totals'].get('analyses_period', 0):,}",
                "Analyses effectuées sur la période",
            ],
            [
                "Analyses (total)",
                f"{kpis['totals'].get('analyses_ever', 0):,}",
                "Total depuis le début",
            ],
            [
                "Mots analysés",
                f"{kpis['totals'].get('words_analyzed', 0):,}",
                "Mots traités au total",
            ],
            [
                "Temps de lecture",
                f"{kpis['totals'].get('reading_time_hours', 0):.1f}h",
                "Heures de contenu analysé",
            ],
        ]

        if kpis.get("averages"):
            kpi_data.extend(
                [
                    [
                        "Complexité moyenne",
                        f"{kpis['averages'].get('complexity_score', 0):.3f}",
                        "Score de difficulté moyen",
                    ],
                    [
                        "Sentiment moyen",
                        f"{kpis['averages'].get('sentiment_polarity', 0):.3f}",
                        "Polarité sentiment moyenne",
                    ],
                    [
                        "Analyses/jour",
                        f"{kpis['averages'].get('analyses_per_day', 0):.1f}",
                        "Moyenne quotidienne",
                    ],
                ]
            )

        kpi_table = Table(kpi_data, colWidths=[2 * inch, 1.5 * inch, 2.5 * inch])
        kpi_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("ALIGN", (1, 1), (1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
                ]
            )
        )

        elements.append(kpi_table)
        elements.append(Spacer(1, 20))

        # Distributions
        if kpis.get("distributions"):
            elements.extend(self._create_distributions_section(kpis["distributions"]))

        return elements

    def _create_distributions_section(self, distributions: dict[str, Any]) -> list:
        """Crée la section des distributions"""
        elements = []

        elements.append(Paragraph("📊 Distributions", self.styles["CustomHeading2"]))
        elements.append(Spacer(1, 12))

        # Sentiment distribution
        if distributions.get("sentiment"):
            sentiment_data = [["Sentiment", "Nombre"]]
            for sentiment, count in distributions["sentiment"].items():
                sentiment_data.append([sentiment.capitalize(), str(count)])

            sentiment_table = Table(sentiment_data, colWidths=[2 * inch, 1 * inch])
            sentiment_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#28a745")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgreen]),
                    ]
                )
            )

            elements.append(Paragraph("Distribution des Sentiments:", self.styles["Normal"]))
            elements.append(sentiment_table)
            elements.append(Spacer(1, 12))

        # Analysis modes distribution
        if distributions.get("analysis_mode"):
            mode_data = [["Mode d'Analyse", "Utilisation"]]
            for mode, count in distributions["analysis_mode"].items():
                mode_data.append([mode.capitalize(), str(count)])

            mode_table = Table(mode_data, colWidths=[2 * inch, 1 * inch])
            mode_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightblue]),
                    ]
                )
            )

            elements.append(Paragraph("Distribution des Modes d'Analyse:", self.styles["Normal"]))
            elements.append(mode_table)
            elements.append(Spacer(1, 12))

        return elements

    def _create_trends_section(self, trends: dict[str, Any]) -> list:
        """Crée la section des tendances"""
        elements = []

        elements.append(Paragraph("📈 Tendances Temporelles", self.styles["CustomHeading2"]))
        elements.append(Spacer(1, 12))

        if not trends or not trends.get("daily_analyses"):
            elements.append(
                Paragraph("Aucune donnée de tendance disponible", self.styles["Normal"])
            )
            return elements

        # daily_analyses is guaranteed non-empty by the guard above.
        daily_analyses = trends["daily_analyses"]
        total_days = len(daily_analyses)
        total_analyses = sum(day["count"] for day in daily_analyses)
        avg_per_day = total_analyses / total_days
        max_day = max(daily_analyses, key=lambda x: x["count"])

        summary_text = f"""
        Résumé des tendances sur {total_days} jours:
        • Total analyses: {total_analyses}
        • Moyenne par jour: {avg_per_day:.1f}
        • Jour le plus actif: {max_day.get("date", "N/A")} ({max_day.get("count", 0)} analyses)
        """

        elements.append(Paragraph(summary_text, self.styles["Normal"]))
        elements.append(Spacer(1, 12))

        # Tableau des analyses quotidiennes (derniers 10 jours)
        recent_data = [["Date", "Analyses"]]
        for day in daily_analyses[-10:]:  # 10 derniers jours
            date_str = datetime.fromisoformat(day["date"]).strftime("%d/%m/%Y")
            recent_data.append([date_str, str(day["count"])])

        trend_table = Table(recent_data, colWidths=[1.5 * inch, 1 * inch])
        trend_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#764ba2")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightpink]),
                ]
            )
        )

        elements.append(Paragraph("Activité des 10 derniers jours:", self.styles["Normal"]))
        elements.append(trend_table)
        elements.append(Spacer(1, 20))

        return elements

    def _create_insights_section(self, insights: dict[str, Any]) -> list:
        """Crée la section des insights"""
        elements = []

        elements.append(
            Paragraph("🧠 Insights et Métriques Avancées", self.styles["CustomHeading2"])
        )
        elements.append(Spacer(1, 12))

        if not insights:
            elements.append(Paragraph("Aucune donnée d'insight disponible", self.styles["Normal"]))
            return elements

        # Content Metrics
        if insights.get("content_metrics"):
            metrics = insights["content_metrics"]

            content_data = [
                ["Métrique", "Valeur"],
                ["Mots moyens par analyse", f"{metrics.get('avg_words_per_analysis', 0):,.0f}"],
                ["Mots uniques moyens", f"{metrics.get('avg_unique_words', 0):,.0f}"],
                ["Richesse vocabulaire", f"{metrics.get('avg_vocabulary_richness', 0):.3f}"],
                ["Complexité moyenne", f"{metrics.get('avg_complexity', 0):.3f}"],
                ["Sentiment moyen", f"{metrics.get('avg_sentiment', 0):.3f}"],
            ]

            content_table = Table(content_data, colWidths=[3 * inch, 1.5 * inch])
            content_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f39c12")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("ALIGN", (1, 1), (1, -1), "CENTER"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightyellow]),
                    ]
                )
            )

            elements.append(Paragraph("Métriques de Contenu:", self.styles["Normal"]))
            elements.append(content_table)
            elements.append(Spacer(1, 15))

        # Top Analyses
        if insights.get("top_analyses"):
            elements.extend(self._create_top_analyses_section(insights["top_analyses"]))

        return elements

    def _create_top_analyses_section(self, top_analyses: dict[str, Any]) -> list:
        """Crée la section des meilleures analyses"""
        elements = []

        elements.append(Paragraph("🏆 Analyses Remarquables", self.styles["Normal"]))
        elements.append(Spacer(1, 8))

        # Most words
        if top_analyses.get("most_words"):
            top_words = top_analyses["most_words"][0] if top_analyses["most_words"] else None
            if top_words:
                elements.append(
                    Paragraph(
                        f"• Plus de mots: {top_words['words']:,} mots (Analyse #{top_words['id']})",
                        self.styles["Normal"],
                    )
                )

        # Most complex
        if top_analyses.get("most_complex"):
            top_complex = top_analyses["most_complex"][0] if top_analyses["most_complex"] else None
            if top_complex:
                elements.append(
                    Paragraph(
                        f"• Plus complexe: {top_complex['complexity']:.3f} (Analyse #{top_complex['id']})",
                        self.styles["Normal"],
                    )
                )

        # Richest vocabulary
        if top_analyses.get("richest_vocabulary"):
            top_vocab = (
                top_analyses["richest_vocabulary"][0]
                if top_analyses["richest_vocabulary"]
                else None
            )
            if top_vocab:
                elements.append(
                    Paragraph(
                        f"• Vocabulaire le plus riche: {top_vocab['richness']:.3f} (Analyse #{top_vocab['id']})",
                        self.styles["Normal"],
                    )
                )

        elements.append(Spacer(1, 15))
        return elements

    def _generate_excel_report(self, data: dict[str, Any], days: int) -> io.BytesIO:
        """Génère un rapport Excel"""
        try:
            import openpyxl

            buffer = io.BytesIO()
            wb = openpyxl.Workbook()

            # Feuille KPIs
            ws_kpis = wb.create_sheet("KPIs")
            del wb["Sheet"]
            ws_kpis.title = "KPIs"
            self._create_excel_kpis_sheet(ws_kpis, data.get("kpis", {}))

            # Feuille Tendances
            ws_trends = wb.create_sheet("Tendances")
            self._create_excel_trends_sheet(ws_trends, data.get("trends", {}))

            # Feuille Insights
            ws_insights = wb.create_sheet("Insights")
            self._create_excel_insights_sheet(ws_insights, data.get("insights", {}))

            wb.save(buffer)
            buffer.seek(0)
            return buffer

        except ImportError:
            logger.error("openpyxl non disponible pour l'export Excel")
            raise ValueError("Export Excel non disponible - bibliothèque manquante")

    def _create_excel_kpis_sheet(self, ws: Worksheet, kpis: dict[str, Any]) -> None:
        """Crée la feuille KPIs Excel"""
        # En-tête
        ws["A1"] = "YouTube Transcript Analyzer - KPIs"
        ws["A1"].font = Font(size=16, bold=True)
        ws["A1"].fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")

        # KPIs principaux
        row = 3
        if kpis.get("totals"):
            ws[f"A{row}"] = "Analyses (période)"
            ws[f"B{row}"] = kpis["totals"].get("analyses_period", 0)
            row += 1

            ws[f"A{row}"] = "Mots analysés"
            ws[f"B{row}"] = kpis["totals"].get("words_analyzed", 0)
            row += 1

            ws[f"A{row}"] = "Temps de lecture (h)"
            ws[f"B{row}"] = kpis["totals"].get("reading_time_hours", 0)
            row += 1

    def _create_excel_trends_sheet(self, ws: Worksheet, trends: dict[str, Any]) -> None:
        """Crée la feuille Tendances Excel"""
        ws["A1"] = "Tendances Temporelles"
        ws["A1"].font = Font(size=16, bold=True)

        if trends.get("daily_analyses"):
            row = 3
            ws[f"A{row}"] = "Date"
            ws[f"B{row}"] = "Analyses"

            for day in trends["daily_analyses"]:
                row += 1
                ws[f"A{row}"] = day["date"]
                ws[f"B{row}"] = day["count"]

    def _create_excel_insights_sheet(self, ws: Worksheet, insights: dict[str, Any]) -> None:
        """Crée la feuille Insights Excel"""
        ws["A1"] = "Insights et Métriques"
        ws["A1"].font = Font(size=16, bold=True)

        if insights.get("content_metrics"):
            row = 3
            metrics = insights["content_metrics"]

            ws[f"A{row}"] = "Mots moyens/analyse"
            ws[f"B{row}"] = metrics.get("avg_words_per_analysis", 0)
            row += 1

            ws[f"A{row}"] = "Complexité moyenne"
            ws[f"B{row}"] = metrics.get("avg_complexity", 0)
            row += 1


# Instance globale du service
report_export_service = ReportExportService()
